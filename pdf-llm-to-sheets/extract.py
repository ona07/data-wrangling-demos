"""レイアウトがそろっていない領収書・請求書の PDF を、1 行 = 1 枚の表にする。

PDF から文字を取り出し（pdfplumber）、Claude に項目を JSON で返させ、**こちらで検算する**。
LLM は読めなかった欄を埋めたがるので、返ってきた数字が元の紙に実在するかまで確かめる。

    python extract.py samples/*.pdf -o out/expenses.xlsx            # Claude に聞く
    python extract.py samples/*.pdf -o out/expenses.xlsx --offline  # 保存済みの応答で動かす（鍵なし）
    python extract.py samples/*.pdf -o out/expenses.xlsx --save-cache  # 応答を cached/ に残す

出力は 3 つ: Excel（一覧 + 明細）、Google Sheets に貼れる TSV、生の JSON。
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).parent
CACHE = HERE / "cached"

# 既定は Claude Opus 5。抽出は難しくないので effort は low で足りる（README に 1 枚あたりの値段）
DEFAULT_MODEL = "claude-opus-5"

SYSTEM = """You extract bookkeeping fields from receipts and invoices.

The documents are photographed or scanned business papers from Japan and abroad, so the
layout, the language (Japanese or English) and the wording of each field all vary.

Rules:
- Copy values exactly as they appear. Do not convert currencies and do not recompute totals.
- Numbers must be plain numbers: 4218, not "4,218" or "¥4,218".
- Dates must be YYYY-MM-DD. Japanese era years (令和 8 = 2026) must be converted.
- If a field is not written on the document, return null. Never guess, never infer a
  plausible value, never carry a number over from another field.
- `notes` is for anything a bookkeeper would need that has no field of its own
  (payment terms, "tax included", zero-rated VAT, the receipt number).
"""

SCHEMA = {
    "type": "object",
    "properties": {
        "doc_type": {"type": "string",
                     "description": "receipt / invoice / expense_report, as written"},
        "issuer": {"type": ["string", "null"], "description": "who is paid"},
        "date": {"type": ["string", "null"], "description": "YYYY-MM-DD"},
        "currency": {"type": "string", "description": "JPY / USD / GBP ..."},
        "registration_no": {"type": ["string", "null"],
                            "description": "invoice registration or VAT number"},
        "subtotal": {"type": ["number", "null"]},
        "tax": {"type": ["number", "null"]},
        "total": {"type": ["number", "null"]},
        "items": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "description": {"type": "string"},
                    "qty": {"type": ["number", "null"]},
                    "unit_price": {"type": ["number", "null"]},
                    "amount": {"type": ["number", "null"]},
                },
                "required": ["description", "qty", "unit_price", "amount"],
                "additionalProperties": False,
            },
        },
        "notes": {"type": ["string", "null"]},
    },
    "required": ["doc_type", "issuer", "date", "currency", "registration_no",
                 "subtotal", "tax", "total", "items", "notes"],
    "additionalProperties": False,
}


def read_text(path: Path) -> str:
    """PDF の全ページの文字。画像だけの PDF ならここが空になる（README 参照）。"""
    with pdfplumber.open(path) as pdf:
        return "\n".join((page.extract_text() or "") for page in pdf.pages).strip()


KEYS_SENTENCE = ("Return ONLY one JSON object, no prose and no code fence, with exactly these keys: "
                 + ", ".join(SCHEMA["required"])
                 + ". Each item of `items` has description, qty, unit_price, amount. "
                   "Use null for anything the document does not state.")


def ask_claude(text: str, model: str, use_schema: bool = True) -> dict:
    """Claude に項目を JSON で返させる。

    既定は structured outputs（`output_config.format`）で形を縛る。返事が必ずその形になるので、
    呼ぶ側に「JSON じゃなかったとき」の分岐が要らない。
    スキーマを受け付けない構成（古い SDK、別の経路）では `--no-schema` で言葉で頼む経路に落とす。
    """
    import anthropic

    client = anthropic.Anthropic()
    prompt = f"Extract the fields from this document.\n\n---\n{text}\n---"
    output_config: dict = {"effort": "low"}       # 抽出は難しくない。考えさせすぎない
    system = SYSTEM
    if use_schema:
        output_config["format"] = {"type": "json_schema", "schema": SCHEMA}
    else:
        system = SYSTEM + "\n" + KEYS_SENTENCE
    try:
        response = client.messages.create(
            model=model, max_tokens=4096, system=system,
            messages=[{"role": "user", "content": prompt}],
            output_config=output_config,
        )
    except anthropic.BadRequestError as error:
        if use_schema:
            raise SystemExit(
                f"スキーマが受け付けられなかった: {error}\n"
                "--no-schema を付けると、同じ項目を言葉で頼む経路で動く。") from error
        raise
    body = next(block.text for block in response.content if block.type == "text")
    return json.loads(re.sub(r"^```(?:json)?|```$", "", body.strip(), flags=re.MULTILINE))


# --------------------------------------------------------------------------
# 検算（ここが売り物。LLM の出力をそのまま信じない）
# --------------------------------------------------------------------------

def numbers_in(text: str) -> set[float]:
    """紙に実際に書かれている数。カンマと通貨記号は外す。"""
    found = set()
    for raw in re.findall(r"\d[\d,]*(?:\.\d+)?", text):
        try:
            found.add(float(raw.replace(",", "")))
        except ValueError:
            continue
    return found


def verify(data: dict, text: str) -> list[str]:
    """合わないところを日本語で並べて返す。空なら OK。"""
    problems: list[str] = []
    on_paper = numbers_in(text)

    # 1. 返ってきた金額が、紙に実在するか（LLM が計算してしまっていないか）
    for field in ("subtotal", "tax", "total"):
        value = data.get(field)
        if value is None:
            continue
        if float(value) not in on_paper:
            problems.append(f"{field} の {value:,} が紙に見つからない")

    # 2. 明細の合計と、小計（無ければ合計）が合うか
    amounts = [i["amount"] for i in data.get("items", []) if i.get("amount") is not None]
    subtotal = data.get("subtotal")
    against, label = (subtotal, "小計") if subtotal is not None else (data.get("total"), "合計")
    if amounts and against is not None and abs(sum(amounts) - float(against)) > 0.5:
        problems.append(
            f"明細の合計 {sum(amounts):,.0f} と{label} {float(against):,.0f} が違う")

    # 3. 小計 + 税 = 合計 か
    tax, total = data.get("tax"), data.get("total")
    if (subtotal is not None and tax is not None and total is not None
            and abs(float(subtotal) + float(tax) - float(total)) > 0.5):
        problems.append(
            f"小計 {float(subtotal):,.0f} + 税 {float(tax):,.0f} が合計 {float(total):,.0f} に合わない")

    # 4. 日付の形
    date = data.get("date")
    if date and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(date)):
        problems.append(f"日付が YYYY-MM-DD でない（{date}）")

    # 5. 支払先が無い（経理が困る）
    if not data.get("issuer"):
        problems.append("支払先を読めていない")
    return problems


# --------------------------------------------------------------------------
# 出力
# --------------------------------------------------------------------------

SUMMARY_COLUMNS = ("ファイル", "種類", "日付", "支払先", "通貨", "小計", "税", "合計",
                   "登録番号", "明細数", "検算", "備考")


def summary_row(name: str, data: dict, problems: list[str]) -> list:
    return [name, data.get("doc_type", ""), data.get("date") or "", data.get("issuer") or "",
            data.get("currency", ""), data.get("subtotal"), data.get("tax"), data.get("total"),
            data.get("registration_no") or "", len(data.get("items", [])),
            "OK" if not problems else "要確認: " + " / ".join(problems),
            data.get("notes") or ""]


def write_excel(rows: list[tuple[str, dict, list[str]]], out: Path) -> None:
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="305496")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")

    def header(ws, columns):
        ws.append(list(columns))
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "一覧"
    header(ws, SUMMARY_COLUMNS)
    for name, data, problems in rows:
        ws.append(summary_row(name, data, problems))
        if problems:
            for cell in ws[ws.max_row]:
                cell.fill = warn_fill
    for column, width in zip("ABCDEFGHIJKL", (26, 14, 12, 28, 7, 11, 9, 11, 18, 7, 46, 34)):
        ws.column_dimensions[column].width = width

    detail = wb.create_sheet("明細")
    header(detail, ("ファイル", "品目", "数量", "単価", "金額"))
    for name, data, _ in rows:
        for item in data.get("items", []):
            detail.append([name, item.get("description", ""), item.get("qty"),
                           item.get("unit_price"), item.get("amount")])
    for column, width in zip("ABCDE", (26, 40, 8, 11, 12)):
        detail.column_dimensions[column].width = width
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def write_tsv(rows: list[tuple[str, dict, list[str]]], out: Path) -> None:
    """Google Sheets にそのまま貼れる形。空欄は空文字。"""
    lines = ["\t".join(SUMMARY_COLUMNS)]
    for name, data, problems in rows:
        lines.append("\t".join("" if v is None else str(v) for v in summary_row(name, data, problems)))
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("pdfs", nargs="+", type=Path)
    parser.add_argument("-o", "--out", type=Path, default=HERE / "out/expenses.xlsx")
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument("--offline", action="store_true",
                        help="cached/ に保存した応答を使う（API の鍵が無くても動く）")
    parser.add_argument("--save-cache", action="store_true", help="応答を cached/ に残す")
    parser.add_argument("--no-schema", action="store_true",
                        help="structured outputs を使わず、項目を言葉で頼む")
    args = parser.parse_args()

    rows: list[tuple[str, dict, list[str]]] = []
    for path in args.pdfs:
        text = read_text(path)
        if not text:
            print(f"× {path.name}: 文字が入っていない PDF（画像だけ）。OCR が要る")
            continue
        cache_file = CACHE / f"{path.stem}.json"
        if args.offline:
            if not cache_file.is_file():
                print(f"× {path.name}: 保存済みの応答が無い（{cache_file}）")
                continue
            data = json.loads(cache_file.read_text(encoding="utf-8"))
        else:
            data = ask_claude(text, args.model, use_schema=not args.no_schema)
            if args.save_cache:
                CACHE.mkdir(exist_ok=True)
                cache_file.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n",
                                      encoding="utf-8")
        problems = verify(data, text)
        rows.append((path.name, data, problems))
        mark = "OK  " if not problems else "要確認"
        print(f"{mark} {path.name}: {data.get('issuer') or '支払先不明'} "
              f"{data.get('currency', '')} {data.get('total')}")
        for problem in problems:
            print(f"       - {problem}")

    if not rows:
        print("読めた PDF が無い")
        return 1
    write_excel(rows, args.out)
    write_tsv(rows, args.out.with_suffix(".tsv"))
    args.out.with_suffix(".json").write_text(
        json.dumps([{"file": n, "data": d, "problems": p} for n, d, p in rows],
                   ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    ok = sum(1 for _, _, p in rows if not p)
    print(f"\n{len(rows)} 枚（検算 OK {ok} 枚 / 要確認 {len(rows) - ok} 枚）")
    print(f"  {args.out}\n  {args.out.with_suffix('.tsv')}\n  {args.out.with_suffix('.json')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
